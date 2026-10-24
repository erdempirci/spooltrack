"""Generates the Excel upload templates (public/templates) and an E2E test data set (test-data).

Run:  python3 scripts/make_templates.py
Requires: openpyxl
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
TPL = ROOT / "public" / "templates"
TEST = ROOT / "test-data"

HEAD_FILL = PatternFill("solid", fgColor="1D5FD1")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def write(path: Path, sheet: str, header, rows, notes=None, validations=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(header, start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        ws.column_dimensions[c.column_letter].width = max(14, len(h) + 4)
    ws.freeze_panes = "A2"
    # Revision column as text so "0" or "01" revisions are not turned into numbers by Excel.
    if "revizyon" in header:
        col = header.index("revizyon") + 1
        for r in range(2, 1001):
            ws.cell(row=r, column=col).number_format = "@"
    for col, values in (validations or {}).items():
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=False)
        dv.add(f"{col}2:{col}5000")
        ws.add_data_validation(dv)
    if notes:
        n = wb.create_sheet("Açıklama")
        for line in notes:
            n.append([line])
        n.column_dimensions["A"].width = 110
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


ISO_H = ["izometri_no", "revizyon", "sistem", "konum_turu", "konum_adi"]
SPOOL_H = ["spool_no", "izometri_no", "revizyon", "baslangic_durumu"]
MAT_H = ["malzeme_kodu", "malzeme_aciklamasi", "cap_olcu", "malzeme_kalitesi", "birim", "toplam_miktar", "depo_stogu"]
SM_H = ["spool_no", "izometri_no", "revizyon", "malzeme_kodu", "miktar"]

COMMON = [
    "Veriler yalnızca İLK sayfadan okunur. Sütun başlıklarını değiştirmeyin; sütun sırası önemli değildir.",
    "Uygulama yüklemeden önce önizleme gösterir: eksik sütun, geçersiz satır, mükerrer ve eşleşmeyen kayıtlar listelenir.",
    "Örnek satırları silip kendi verinizi girin.",
]

# ───────────── Templates (spec examples) ─────────────
write(TPL / "IZOMETRI_SABLONU.xlsx", "Izometriler", ISO_H,
      [["FO-120-001", "A", "Fuel Oil", "Zone", "Z01"],
       ["FO-120-002", "B", "Fuel Oil", "Block", "B12"],
       ["CHW-340-021", "0", "Chilled Water", "Zone", "Z03"]],
      COMMON + [
          "konum_turu: Zone veya Block (zone, ZONE, zon / block, Blok, BLK kabul edilir). Başka bir değer hata verir.",
          "Benzersiz kayıt: izometri_no + revizyon. Aynı izometri ve revizyon tekrar yüklenirse ikinci kayıt oluşmaz.",
          "Yeni revizyon ayrı bir izometri kaydıdır; eski revizyonun spoolları yeni revizyona taşınmaz.",
      ], {"D": ["Zone", "Block"]})

write(TPL / "SPOOL_SABLONU.xlsx", "Spoollar", SPOOL_H,
      [["FO-120-001-SP01", "FO-120-001", "A", "Başlamadı"],
       ["FO-120-001-SP02", "FO-120-001", "A", "Başlamadı"],
       ["FO-120-002-SP01", "FO-120-002", "B", "Devam Ediyor"]],
      COMMON + [
          "Spool, izometri_no + revizyon ile bağlanır. İzometri listesinde olmayan spoollar alınmaz, hata listesinde gösterilir.",
          "baslangic_durumu: Başlamadı / Devam Ediyor / Bitti (Bekliyor, Üretimde, Tamamlandı, Not Started, In Progress, Completed, Done da kabul edilir).",
          "Kayıtlı bir spool tekrar yüklenirse sahadaki güncel durumu korunur.",
      ], {"D": ["Başlamadı", "Devam Ediyor", "Bitti"]})

write(TPL / "TOPLAM_MALZEME_SABLONU.xlsx", "Malzemeler", MAT_H,
      [["PIPE-DN50-CS", "Boru", "DN50", "Carbon Steel", "m", 240, 190],
       ["ELB-DN50-CS", "90 derece dirsek", "DN50", "Carbon Steel", "adet", 48, 20],
       ["FLG-DN50-CS", "Flanş", "DN50", "Carbon Steel", "adet", 36, 18]],
      COMMON + [
          "Benzersiz kayıt: malzeme_kodu. Aynı kod tekrar yüklenirse kayıt güncellenir (ör. depo stoğu).",
          "Miktarlarda ondalık ayırıcı olarak virgül veya nokta kullanılabilir.",
      ])

write(TPL / "SPOOL_MALZEME_SABLONU.xlsx", "SpoolMalzeme", SM_H,
      [["FO-120-001-SP01", "FO-120-001", "A", "PIPE-DN50-CS", 8],
       ["FO-120-001-SP01", "FO-120-001", "A", "ELB-DN50-CS", 3],
       ["FO-120-001-SP01", "FO-120-001", "A", "FLG-DN50-CS", 2]],
      COMMON + [
          "Spool, izometri, revizyon ve malzeme kodu sistemde kayıtlı olmalıdır; aksi hâlde satır alınmaz.",
          "Aynı spool + malzeme kodu bir dosyada bir kez yazılmalıdır (miktarları tek satırda toplayın).",
      ])

# ───────────── E2E test data (realistic-sized, with deliberate errors) ─────────────
isos, spools, sm = [], [], []
systems = [("Fuel Oil", "FO"), ("Chilled Water", "CHW"), ("Sea Water", "SW")]
locs = [("Zone", "Z01"), ("Zone", "Z02"), ("Block", "B12")]
n = 0
for (lt, ln) in locs:
    for (sysname, pre) in systems:
        for k in range(1, 4):
            n += 1
            iso = f"{pre}-{100 + n:03d}-{k:03d}"
            isos.append([iso, "A", sysname, lt, ln])
            for s in range(1, 4):
                sp = f"{iso}-SP{s:02d}"
                spools.append([sp, iso, "A", "Başlamadı"])
                sm.append([sp, iso, "A", "PIPE-DN50-CS", 2.5])
                sm.append([sp, iso, "A", "ELB-DN50-CS", 1])
# Lower-case variants and an iso with no spools
isos.append(["FO-999-001", "A", "fuel oil", "zon", "z01"])          # normalised → Fuel Oil / Zone / Z01
# Deliberate errors
isos.append(["BAD-001", "A", "Fuel Oil", "Deck", "D1"])            # unknown konum_turu
isos.append(["", "A", "Fuel Oil", "Zone", "Z01"])                  # missing izometri_no
isos.append([isos[0][0], "A", "Fuel Oil", "Zone", "Z01"])          # in-file duplicate
write(TEST / "1_izometri_test.xlsx", "Izometriler", ISO_H, isos)

spools[1][3] = "Devam Ediyor"
spools[2][3] = "Tamamlandı"
spools.append([f"{isos[0][0]}-SP09", isos[0][0], "B", "Başlamadı"])  # wrong revision → unmatched
spools.append(["X-SP01", "NOPE-001", "A", "Başlamadı"])              # unknown iso → unmatched
spools.append([f"{isos[0][0]}-SP10", isos[0][0], "A", "belki"])     # unknown status
write(TEST / "2_spool_test.xlsx", "Spoollar", SPOOL_H, spools)

write(TEST / "3_malzeme_test.xlsx", "Malzemeler", MAT_H,
      [["PIPE-DN50-CS", "Boru", "DN50", "Carbon Steel", "m", 240, 190],
       ["ELB-DN50-CS", "90 derece dirsek", "DN50", "Carbon Steel", "adet", 90, 20],
       ["FLG-DN50-CS", "Flanş", "DN50", "Carbon Steel", "adet", "36", "18"]])

sm.append([spools[0][0], spools[0][1], "A", "UNKNOWN-CODE", 1])  # unknown material
write(TEST / "4_spool_malzeme_test.xlsx", "SpoolMalzeme", SM_H, sm)

write(TEST / "eksik_sutun_test.xlsx", "Izometriler", ["izometri_no", "sistem"], [["A-1", "Fuel Oil"]])
print("templates:", sorted(p.name for p in TPL.glob("*.xlsx")))
print("test data:", sorted(p.name for p in TEST.glob("*.xlsx")), "isos", len(isos), "spools", len(spools), "spool-material", len(sm))
